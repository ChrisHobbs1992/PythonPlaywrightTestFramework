import json
from datetime import datetime, timedelta
from docx import Document
import os

# First run the tests to generate a report
# $ pytest --json-report --json-report-file=report.json

# Then run the report generator code to embed these into microsoft word
# python hooks\\reporting\\generate_docx_report.py


# === CONFIG ===
BASE_DIR = os.path.dirname(os.path.abspath(__file__))
REPORTS_DIR = os.path.join(BASE_DIR, "..", "..", "reports")
JSON_REPORT =  os.path.join(BASE_DIR, "..", "..", "reports", "report.json")
TEMPLATE_DOCX = "template.docx"
OUTPUT_DOCX = os.path.join(REPORTS_DIR, "Test_Report" + "{:%B %d %Y %I %M %p}".format(datetime.now()) + ".docx")

def read_json_report(path):
    if not os.path.exists(path):
        raise FileNotFoundError(f"JSON report not found at {path}")
    with open(path, "r", encoding="utf-8") as f:
        return json.load(f)

def fill_placeholders(doc, placeholders):
    for p in doc.paragraphs:
        for key, val in placeholders.items():
            placeholder = f"{{{{{key}}}}}"  # e.g. {{start_time}}
            if placeholder in p.text:
                inline = p.runs
                for i in range(len(inline)):
                    if placeholder in inline[i].text:
                        inline[i].text = inline[i].text.replace(placeholder, val)

def generate_report():
    print("Loading JSON report...")
    data = read_json_report(JSON_REPORT)

    # Extract data
    created = data["created"]
    if isinstance(created, (int, float)):
        start_time = datetime.fromtimestamp(created)
    else:
        start_time = datetime.fromisoformat(created.replace("Z", ""))

    duration = timedelta(seconds=data["duration"])
    end_time = start_time + duration

    summary = data.get("summary", {})
    total = summary.get("total", 0)
    passed = summary.get("passed", 0)
    failed = summary.get("failed", 0)
    skipped = summary.get("skipped", 0)

    # Prepare placeholders
    placeholders = {
        "start_time": "{:%B %d %Y %I:%M %p}".format(start_time),
        "end_time": "{:%B %d %Y %I:%M %p}".format(end_time),
        "total": str(total),
        "passed": str(passed),
        "failed": str(failed),
        "skipped": str(skipped),
    }

    print("Filling Word template...")
    doc = Document(TEMPLATE_DOCX)
    fill_placeholders(doc, placeholders)
    doc.save(OUTPUT_DOCX)

    print(f"Report generated: {OUTPUT_DOCX}")

if __name__ == "__main__":
    generate_report()
